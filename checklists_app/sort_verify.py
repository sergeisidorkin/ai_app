import io
import json
import logging
import os
import shutil
from datetime import timedelta
from pathlib import Path

from django.conf import settings
from django.db import close_old_connections
from django.db.models import Max
from django.utils import timezone

from core.cloud_storage import CloudStorageNotReadyError, download_file
from core.dsh_run import DshRunError, run_headless

from .models import ChecklistItem, ChecklistSortProposal, ChecklistSortRun
from .sort_parse import (
    dest_code_nn,
    dest_leaf_name,
    finalize_verify_action,
    normalize_rel_path,
    parse_verify_classify_response,
    parse_verify_select_response,
)
from .sort_workspace import (
    _safe_relpath,
    load_cloud_file_sizes,
    resolve_inbox_folder,
    workspace_root_for,
)

logger = logging.getLogger(__name__)

STALE_RUNNING_MESSAGE = "Проверка не завершилась. Нажмите «Проверить» ещё раз."

MAX_PEEK_FILES = 3
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_PEEK_CHARS = 20_000
MAX_PDF_PAGES = 5
SKIP_NAMES = {".ds_store", "plot.log"}
SKIP_SUFFIXES = {".sig", ".p7s", ".bak"}


class VerifyError(Exception):
    pass


class VerifyConflict(Exception):
    pass


def verify_dir_for(run, proposal_id):
    root = Path(run.workspace_path) if (run.workspace_path or "").strip() else workspace_root_for(run.id)
    return root / "verify" / str(proposal_id)


def remove_verify_dir(run, proposal_id):
    path = verify_dir_for(run, proposal_id)
    if path.is_symlink() or path.is_file():
        path.unlink(missing_ok=True)
    elif path.is_dir():
        shutil.rmtree(path, ignore_errors=True)
    parent = path.parent
    try:
        if parent.is_dir() and parent.name == "verify" and not any(parent.iterdir()):
            parent.rmdir()
    except OSError:
        return


def _is_junk_name(name):
    text = str(name or "")
    folded = text.casefold()
    if folded in SKIP_NAMES or text.startswith("."):
        return True
    suffix = Path(text).suffix.casefold()
    return suffix in SKIP_SUFFIXES


def _exists(path: Path):
    return path.exists() or os.path.lexists(path)


def resolve_kit_path(run, kit_path):
    inbox = Path(run.workspace_path or "") / "inbox"
    rel = _safe_relpath(kit_path)
    if rel is None:
        raise VerifyError("Пустой комплект.")
    candidates = [inbox / rel]
    section = _safe_relpath(run.inbox_section_name or "")
    if section is not None:
        candidates.append(inbox / section / rel)
    for candidate in candidates:
        if _exists(candidate):
            return candidate
    raise VerifyError(f"Комплект не найден в workspace: {kit_path}")


def _file_size(path: Path, cloud_root=None, cloud_sizes=None):
    if cloud_root is not None and cloud_sizes is not None:
        try:
            key = path.relative_to(cloud_root).as_posix()
        except ValueError:
            key = ""
        if key in cloud_sizes:
            return cloud_sizes[key]
    try:
        return int(path.stat().st_size)
    except OSError:
        return 0


def iter_kit_files(kit_path: Path):
    if kit_path.is_file() or (not kit_path.is_dir() and _exists(kit_path)):
        yield kit_path.name, kit_path
        return
    if not kit_path.is_dir():
        return
    for dirpath, dirnames, filenames in os.walk(kit_path):
        dirnames[:] = [name for name in dirnames if not name.startswith(".")]
        for name in filenames:
            full = Path(dirpath) / name
            rel = full.relative_to(kit_path).as_posix()
            yield rel, full


def build_kit_manifest(kit_path: Path):
    rows = []
    cloud_root, cloud_sizes = load_cloud_file_sizes(kit_path)
    for rel, full in iter_kit_files(kit_path):
        name = Path(rel).name
        if _is_junk_name(name):
            continue
        size = _file_size(full, cloud_root=cloud_root, cloud_sizes=cloud_sizes)
        skipped = size > MAX_FILE_BYTES
        rows.append({
            "path": rel.replace("\\", "/"),
            "size": size,
            "ext": Path(rel).suffix.casefold(),
            "skipped": skipped,
        })
    rows.sort(key=lambda row: row["path"].casefold())
    return rows


def readable_manifest_paths(manifest):
    return [row["path"] for row in manifest if not row.get("skipped")]


def accept_peek_paths(requested, manifest, *, kit_prefix=""):
    by_key = {}
    for row in manifest:
        rel = normalize_rel_path(row.get("path") or "", kit_prefix=kit_prefix)
        if rel:
            by_key[rel.casefold()] = row
    accepted = []
    seen = set()
    for raw in requested or []:
        rel = normalize_rel_path(raw, kit_prefix=kit_prefix)
        if not rel:
            continue
        row = by_key.get(rel.casefold())
        if row is None or row.get("skipped"):
            continue
        key = rel.casefold()
        if key in seen:
            continue
        seen.add(key)
        accepted.append(row["path"])
        if len(accepted) >= MAX_PEEK_FILES:
            break
    return accepted


def extract_text_from_bytes(filename, data, *, budget):
    if budget <= 0:
        return ""
    name = str(filename or "").casefold()
    payload = data or b""
    if name.endswith(".pdf"):
        text = _extract_pdf(payload)
    elif name.endswith(".docx"):
        text = _extract_docx(payload)
    elif name.endswith((".txt", ".md", ".csv", ".log")):
        text = payload.decode("utf-8", errors="replace")
    else:
        return ""
    text = " ".join(str(text or "").split())
    return text[:budget]


def _extract_pdf(data):
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception:
        return ""
    chunks = []
    for page in list(reader.pages)[:MAX_PDF_PAGES]:
        try:
            chunks.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(chunks)


def _extract_docx(data):
    try:
        from docx import Document
    except ImportError:
        return ""
    try:
        document = Document(io.BytesIO(data))
    except Exception:
        return ""
    chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            chunks.append(" ".join(cell.text for cell in row.cells if cell.text))
    return "\n".join(chunks)


def write_peek_markdown(path: Path, kit_path, excerpts):
    lines = [
        f"# Выдержки комплекта {kit_path}",
        "",
        "Читай только этот файл и requests.md. Не открывай бинарники.",
        "",
    ]
    for item in excerpts:
        rel = item["path"]
        text = (item.get("text") or "").strip() or "(текст не извлечён)"
        lines.append(f"## {rel}")
        lines.append("")
        lines.append(text)
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def _cloud_path_for_inbox_rel(run, disk_path, inbox_rel: Path):
    root = str(disk_path or "").rstrip("/")
    if not root:
        return ""
    section = (run.inbox_section_name or "").strip().strip("/")
    parts = [part for part in inbox_rel.parts if part]
    if section and parts and parts[0] == section:
        parts = parts[1:]
    if not parts:
        return root
    return f"{root}/{'/'.join(parts)}"


def _read_local_bytes(path: Path):
    try:
        return path.read_bytes()
    except OSError as exc:
        raise VerifyError(f"Не удалось прочитать файл: {path.name}") from exc


def _download_selected_file(run, user, kit_path: Path, rel, dest: Path):
    source = kit_path / rel if kit_path.is_dir() else kit_path
    cloud_root, cloud_sizes = load_cloud_file_sizes(source)
    size = _file_size(source, cloud_root=cloud_root, cloud_sizes=cloud_sizes)
    if size > MAX_FILE_BYTES:
        raise VerifyError(f"Файл превышает лимит 50 МБ: {rel}")
    if (run.source_kind or "cloud") == "local":
        if source.is_file():
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, dest)
            return dest.read_bytes()
        raise VerifyError(f"Файл не найден: {rel}")

    folder = resolve_inbox_folder(run.project, run.section, run.asset_name)
    if folder is None or not folder.disk_path:
        if source.is_file() and source.stat().st_size:
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, dest)
            return dest.read_bytes()
        raise VerifyError("Нет облачного пути исходных данных для комплекта.")
    kit_rel = _safe_relpath(str(kit_path.relative_to(Path(run.workspace_path) / "inbox")))
    inbox_file = kit_rel / Path(rel) if kit_path.is_dir() else kit_rel
    cloud_path = _cloud_path_for_inbox_rel(run, folder.disk_path, inbox_file)
    try:
        _mime, data = download_file(user, cloud_path)
    except CloudStorageNotReadyError as exc:
        raise VerifyError(str(exc)) from exc
    if not data:
        raise VerifyError(f"Не удалось скачать {rel}.")
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(data)
    return data


def _candidate_lines(run):
    from .sort_service import _dest_section_label, _section_dest_items

    dest_label = _dest_section_label(run.section)
    lines = [f"Типовой раздел dest: {dest_label}."]
    for item in _section_dest_items(run.project_id, run.section_id):
        folder = item.get("folder") or ""
        name = item.get("request_name") or ""
        quote = item.get("quote") or ""
        lines.append(f"- {folder} | {name} | {quote}")
    return "\n".join(lines)


def _select_prompt(run, proposal, files_rel):
    from .sort_service import _dest_section_label

    dest_label = _dest_section_label(run.section)
    return (
        "/checklist-file-verify\n\n"
        "Фаза: select\n"
        f"{_candidate_lines(run)}\n\n"
        f"Комплект: {proposal.kit_path}\n"
        f"Текущая гипотеза dest: {proposal.dest_path or dest_label}\n"
        f"Уверенность: {proposal.confidence or 'не задана'}\n"
        f"Действие сейчас: review\n"
        f"Список файлов: {files_rel}\n"
        "Не скачивай и не читай содержимое. Не ходи вне корня задачи.\n"
        "Укажи до трёх путей из списка, которые отличают пункты dest этого раздела.\n"
        "Выведи JSON-объект: {\"peek\": [\"rel/path\"], \"reason\": \"...\"}.\n"
        "Только пути из списка. Не больше трёх. Пропущенные (skipped) не бери.\n"
    )


def _classify_prompt(run, proposal, peek_rel):
    from .sort_service import _dest_section_label

    dest_label = _dest_section_label(run.section)
    return (
        "/checklist-file-verify\n\n"
        "Фаза: classify\n"
        f"{_candidate_lines(run)}\n\n"
        f"Комплект: {proposal.kit_path}\n"
        f"Текущая гипотеза dest: {proposal.dest_path or dest_label}\n"
        f"Уверенность: {proposal.confidence or 'не задана'}\n"
        f"Выдержки: {peek_rel}\n"
        "Читай только этот peek.md и requests.md. Бинарники не открывай.\n"
        "Режим: только dry-run. Файлы не копировать и не перемещать.\n"
        "После разбора выведи JSON одного объекта с полями "
        "kit, files, dest, name, quote, confidence, action.\n"
        "quote — дословный фрагмент из requests.md, не текст документа.\n"
        "высокая + конкретный пункт dest → move.\n"
        "Если уверенность не стала высокой — action confirm, не review.\n"
        "requests.md содержит все разделы. Если выдержки однозначно соответствуют "
        "другому пункту, в том числе другого раздела — укажи его dest, не держи комплект "
        "в текущей гипотезе.\n"
        "Цитата из документа нужна для себя, в JSON её не клади.\n"
        f"kit оставь {proposal.kit_path}.\n"
    )


def resolve_dest_item(project, dest_path):
    code = dest_code_nn(dest_leaf_name(dest_path))
    if not code:
        return None
    from yandexdisk_app.workspace import _build_item_folder_name

    items = (
        ChecklistItem.objects.filter(project=project)
        .select_related("section")
        .order_by("position", "id")
    )
    for item in items:
        if dest_code_nn(_build_item_folder_name(item)) == code:
            return item
    return None


def canonical_dest_path(project, item):
    from .sort_workspace import section_folder_name
    from yandexdisk_app.workspace import _build_item_folder_name

    return f"dest/{section_folder_name(project, item.section)}/{_build_item_folder_name(item)}"


def ensure_host_run(source_run, target_section, user=None):
    from .sort_service import active_run_for, latest_run_for

    project = source_run.project
    asset = source_run.asset_name or ""
    active = active_run_for(project, target_section, asset)
    if active and active.pk != source_run.pk:
        raise VerifyConflict("Сортировка раздела назначения ещё выполняется.")
    latest = latest_run_for(project, target_section, asset)
    if latest and latest.status == ChecklistSortRun.Status.DONE:
        return latest
    now = timezone.now()
    return ChecklistSortRun.objects.create(
        project=project,
        section=target_section,
        asset_name=asset,
        inbox_section_name=source_run.inbox_section_name,
        source_kind=source_run.source_kind,
        local_inbox_path=source_run.local_inbox_path,
        status=ChecklistSortRun.Status.DONE,
        started_by=user or source_run.started_by,
        workspace_path=source_run.workspace_path,
        started_at=now,
        finished_at=now,
    )


def relocate_proposal_to_item(proposal, item, user=None):
    if item is None or item.section_id == proposal.run.section_id:
        return False
    host = ensure_host_run(proposal.run, item.section, user=user)
    ChecklistSortProposal.objects.filter(
        run=host,
        kit_path=proposal.kit_path,
    ).exclude(pk=proposal.pk).delete()
    max_pos = host.proposals.aggregate(Max("position"))["position__max"] or 0
    proposal.run = host
    proposal.position = max_pos + 1
    proposal.dest_path = canonical_dest_path(proposal.run.project, item)
    return True


def adopt_cross_section_proposals(project, asset_name="", user=None):
    from .sort_parse import normalize_proposal_action

    asset = (asset_name or "").strip()
    proposals = list(
        ChecklistSortProposal.objects.filter(
            run__project=project,
            run__asset_name=asset,
            run__status=ChecklistSortRun.Status.DONE,
        )
        .exclude(kit_path="")
        .select_related("run", "run__project", "run__section")
    )
    changed = 0
    for proposal in proposals:
        item = resolve_dest_item(proposal.run.project, proposal.dest_path)
        if item is None or item.section_id == proposal.run.section_id:
            continue
        try:
            if not relocate_proposal_to_item(proposal, item, user=user):
                continue
        except VerifyConflict:
            continue
        proposal.request_name = proposal.request_name or item.short_name or ""
        proposal.quote = proposal.quote or item.name or ""
        proposal.action = normalize_proposal_action(
            proposal.confidence,
            proposal.dest_path,
            proposal.action,
        )
        proposal.save(update_fields=["run", "position", "dest_path", "request_name", "quote", "action"])
        changed += 1
    return changed


def apply_classify_row(proposal, row, user=None):
    previous_confidence = proposal.confidence
    dest = (row or {}).get("dest_path") or proposal.dest_path
    confidence = (row or {}).get("confidence") or proposal.confidence
    request_name = (row or {}).get("request_name") or proposal.request_name
    quote = (row or {}).get("quote") or proposal.quote
    item = resolve_dest_item(proposal.run.project, dest)
    extra_fields = []
    if item is not None:
        dest = canonical_dest_path(proposal.run.project, item)
        request_name = request_name or item.short_name or ""
        quote = quote or item.name or ""
        if relocate_proposal_to_item(proposal, item, user=user):
            extra_fields.extend(["run", "position"])
            dest = proposal.dest_path
    action = finalize_verify_action(
        previous_confidence,
        confidence,
        dest,
        (row or {}).get("action") or "",
    )
    proposal.dest_path = dest
    proposal.confidence = confidence
    proposal.request_name = request_name
    proposal.quote = quote
    proposal.action = action
    proposal.verify_status = ""
    proposal.verify_error = ""
    proposal.verify_started_at = None
    proposal.save(update_fields=[
        "dest_path",
        "confidence",
        "request_name",
        "quote",
        "action",
        "verify_status",
        "verify_error",
        "verify_started_at",
        *extra_fields,
    ])


def _verify_timeout_seconds():
    return int(getattr(settings, "DSH_VERIFY_TIMEOUT", 180) or 180)


def _verify_stale_seconds():
    return max((_verify_timeout_seconds() * 3) + 300, 900)


def _fail_verify(proposal, message):
    proposal.verify_status = "error"
    proposal.verify_error = str(message)
    proposal.verify_started_at = None
    proposal.save(update_fields=["verify_status", "verify_error", "verify_started_at"])


def reclaim_stale_verifies(run_id=None):
    qs = ChecklistSortProposal.objects.filter(verify_status="running")
    if run_id is not None:
        qs = qs.filter(run_id=run_id)
    rows = list(qs.only("id", "verify_started_at"))
    if not rows:
        return 0
    now = timezone.now()
    stale_limit = timedelta(seconds=_verify_stale_seconds())
    stale_ids = []
    for row in rows:
        started = row.verify_started_at
        if started is None or (now - started) > stale_limit:
            stale_ids.append(row.id)
    if not stale_ids:
        return 0
    return ChecklistSortProposal.objects.filter(pk__in=stale_ids).update(
        verify_status="error",
        verify_error=STALE_RUNNING_MESSAGE,
        verify_started_at=None,
    )


def execute_verify_proposal(proposal_id, user=None, *, close_connections=False):
    if close_connections:
        close_old_connections()
    proposal = None
    run = None
    try:
        claimed = ChecklistSortProposal.objects.filter(
            pk=proposal_id,
            verify_status="queued",
        ).update(
            verify_status="running",
            verify_started_at=timezone.now(),
        )
        if not claimed:
            return False
        proposal = (
            ChecklistSortProposal.objects.select_related(
                "run",
                "run__project",
                "run__section",
                "run__started_by",
                "verify_started_by",
            )
            .filter(pk=proposal_id)
            .first()
        )
        if proposal is None:
            return False
        run = proposal.run
        if user is None:
            user = proposal.verify_started_by or run.started_by
        logger.info("Checklist sort verify %s started kit=%s", proposal_id, proposal.kit_path)
        work = verify_dir_for(run, proposal.id)
        if not (run.workspace_path or "").strip() or not Path(run.workspace_path).exists():
            raise VerifyError("Workspace прогона не найден. Запустите «Распределить» снова.")
        kit_path = resolve_kit_path(run, proposal.kit_path)
        manifest = build_kit_manifest(kit_path)
        readable = readable_manifest_paths(manifest)
        if not readable:
            raise VerifyError("В комплекте нет файлов, из которых можно взять выдержки.")
        work.mkdir(parents=True, exist_ok=True)
        files_path = work / "files.json"
        files_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        files_rel = files_path.relative_to(Path(run.workspace_path)).as_posix()
        timeout = _verify_timeout_seconds()
        if len(readable) == 1:
            selected = readable
        else:
            output = run_headless(
                _select_prompt(run, proposal, files_rel),
                cwd=run.workspace_path,
                timeout=timeout,
            )
            selected = accept_peek_paths(
                parse_verify_select_response(output),
                manifest,
                kit_prefix=str(proposal.kit_path or ""),
            )
            if not selected:
                raise VerifyError("Модель не указала допустимые файлы для проверки.")
        excerpts = []
        budget = MAX_PEEK_CHARS
        copies = work / "files"
        for rel in selected:
            dest = copies / Path(rel)
            data = _download_selected_file(run, user, kit_path, rel, dest)
            text = extract_text_from_bytes(rel, data, budget=budget)
            budget -= len(text)
            excerpts.append({"path": rel, "text": text})
        peek_path = work / "peek.md"
        write_peek_markdown(peek_path, proposal.kit_path, excerpts)
        peek_rel = peek_path.relative_to(Path(run.workspace_path)).as_posix()
        classified = run_headless(
            _classify_prompt(run, proposal, peek_rel),
            cwd=run.workspace_path,
            timeout=timeout,
        )
        row = parse_verify_classify_response(classified, kit_path=proposal.kit_path)
        if row is None:
            raise VerifyError("Не удалось разобрать ответ проверки.")
        row["kit_path"] = proposal.kit_path
        apply_classify_row(proposal, row, user=user)
        proposal.refresh_from_db()
        logger.info("Checklist sort verify %s finished action=%s", proposal_id, proposal.action)
        return True
    except (VerifyError, DshRunError) as exc:
        if proposal is not None:
            _fail_verify(proposal, exc)
        logger.warning("Checklist sort verify %s failed: %s", proposal_id, exc)
    except Exception as exc:
        logger.exception("Checklist sort verify %s failed", proposal_id)
        if proposal is not None:
            _fail_verify(proposal, exc)
    finally:
        if run is not None:
            remove_verify_dir(run, proposal_id)
        if close_connections:
            close_old_connections()


def start_verify_proposal(*, proposal, user):
    run = proposal.run
    if run.status in {ChecklistSortRun.Status.QUEUED, ChecklistSortRun.Status.RUNNING}:
        raise VerifyConflict("Сортировка этого раздела ещё выполняется.")
    if run.status != ChecklistSortRun.Status.DONE:
        raise VerifyError("Проверить можно только после завершённого «Распределить».")
    if not (proposal.kit_path or "").strip():
        raise VerifyError("Нет комплекта для проверки.")
    if (proposal.action or "").lower() != "review":
        raise VerifyError("Проверка доступна только для строк «Проверить».")
    reclaim_stale_verifies(run_id=run.pk)
    proposal.refresh_from_db()
    if proposal.verify_status in {"queued", "running"}:
        raise VerifyConflict("Эта строка уже проверяется.")
    from .sort_service import active_run_for

    active = active_run_for(run.project, run.section, run.asset_name)
    if active and active.pk != run.pk:
        raise VerifyConflict("Сортировка этого раздела уже выполняется.")
    if not (getattr(settings, "DSH_HEADLESS_CMD", "") or "").strip():
        raise VerifyError(
            "DSH для сортировки не настроен. Локально запустите ./scripts/dev_dsh.sh; "
            "на проде задайте DSH_HEADLESS_CMD из deploy/dsh/prod.env.dsh.example."
        )

    proposal.verify_status = "queued"
    proposal.verify_error = ""
    proposal.verify_started_at = None
    proposal.verify_started_by = user
    proposal.save(update_fields=[
        "verify_status",
        "verify_error",
        "verify_started_at",
        "verify_started_by",
    ])

    inline = bool(getattr(settings, "DSH_SORT_INLINE", False))
    if inline:
        execute_verify_proposal(proposal.id, user=user)
        proposal.refresh_from_db()
        return proposal

    return proposal
