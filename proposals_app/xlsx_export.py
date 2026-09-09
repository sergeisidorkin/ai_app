import io
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


TRAVEL_LABELS = {"Командировочные расходы", "Командировочные расходы, евро"}
MAX_STAGES = 20
MAX_ASSETS = 50
MAX_ROWS_PER_BLOCK = 500


def _text(value, limit=1000):
    return str(value or "").strip()[:limit]


def _excel_text(value):
    text = str(value or "")
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


def _number(value):
    raw = _text(value).replace("\xa0", "").replace(" ", "").replace(",", ".").replace("%", "")
    if not raw:
        return None
    try:
        return float(Decimal(raw))
    except (InvalidOperation, ValueError, TypeError):
        return None


def _is_true(value):
    if value is True:
        return True
    return _text(value).lower() in {"1", "true", "yes", "on"}


def _is_travel(row):
    return _text(row.get("service_name")) in TRAVEL_LABELS


def _group_key(row):
    code = "" if _is_true(row.get("merge_without_code")) else _text(row.get("code"), 100)
    return (
        _text(row.get("specialist"), 255),
        _text(row.get("job_title"), 255),
        code,
    )


def _formula_sum(cells):
    refs = [str(cell) for cell in cells if cell]
    if not refs:
        return ""
    if len(refs) == 1:
        return refs[0]
    return f"SUM({','.join(refs)})"


def _formula_nonzero_sum(cells):
    expression = _formula_sum(cells)
    return f'=IF({expression}=0,"",{expression})' if expression else ""


def _normalize_rows(value, asset_count):
    if not isinstance(value, list):
        raise ValueError("Строки коммерческого предложения переданы в некорректном формате.")
    if len(value) > MAX_ROWS_PER_BLOCK:
        raise ValueError("Слишком много строк коммерческого предложения.")
    rows = []
    for item in value:
        if not isinstance(item, dict):
            continue
        day_counts = item.get("asset_day_counts")
        if not isinstance(day_counts, list):
            day_counts = []
        rows.append(
            {
                "specialist": _text(item.get("specialist"), 255),
                "job_title": _text(item.get("job_title"), 255),
                "professional_status": _text(item.get("professional_status"), 255),
                "code": _text(item.get("code"), 100),
                "service_name": _text(item.get("service_name"), 500),
                "merge_without_code": _is_true(item.get("merge_without_code")),
                "rate_eur_per_day": _number(item.get("rate_eur_per_day")),
                "asset_day_counts": [
                    _number(day_counts[index]) if index < len(day_counts) else None
                    for index in range(asset_count)
                ],
                "total_eur_without_vat": _number(item.get("total_eur_without_vat")),
            }
        )
    return rows


def normalize_snapshot(payload):
    if not isinstance(payload, dict):
        raise ValueError("Данные для экспорта переданы в некорректном формате.")
    assets_raw = payload.get("assets")
    if not isinstance(assets_raw, list):
        assets_raw = []
    if len(assets_raw) > MAX_ASSETS:
        raise ValueError("Слишком много активов для экспорта.")
    assets = [_text(value, 255) or f"Актив {index + 1}" for index, value in enumerate(assets_raw)]
    if not assets:
        assets = ["Актив 1"]

    stages_raw = payload.get("stages")
    if not isinstance(stages_raw, list) or not stages_raw:
        raise ValueError("Нет этапов для экспорта.")
    if len(stages_raw) > MAX_STAGES:
        raise ValueError("Слишком много этапов для экспорта.")

    stages = []
    for index, item in enumerate(stages_raw):
        if not isinstance(item, dict):
            raise ValueError("Этап передан в некорректном формате.")
        totals = item.get("totals") if isinstance(item.get("totals"), dict) else {}
        stages.append(
            {
                "label": _text(item.get("label"), 500) or f"Коммерческое предложение: Этап {index + 1}",
                "product_label": _text(item.get("product_label"), 255),
                "rows": _normalize_rows(item.get("rows"), len(assets)),
                "totals": {
                    "exchange_rate": _number(totals.get("exchange_rate")),
                    "discount_percent": _number(totals.get("discount_percent")),
                    "contract_total": _number(totals.get("contract_total")),
                    "contract_total_auto": _number(totals.get("contract_total_auto")),
                    "rub_total_service_text": _text(totals.get("rub_total_service_text"), 500),
                    "discounted_total_service_text": _text(totals.get("discounted_total_service_text"), 500),
                    "travel_expenses_mode": _text(totals.get("travel_expenses_mode"), 20) or "actual",
                },
            }
        )

    summary = None
    summary_raw = payload.get("summary")
    if len(stages) > 1 and not isinstance(summary_raw, dict):
        raise ValueError("Нет сводного блока коммерческого предложения.")
    if len(stages) > 1 and isinstance(summary_raw, dict):
        totals = summary_raw.get("totals") if isinstance(summary_raw.get("totals"), dict) else {}
        summary = {
            "label": _text(summary_raw.get("label"), 500) or "Коммерческое предложение: все этапы",
            "rows": _normalize_rows(summary_raw.get("rows"), len(assets)),
            "totals": {
                "exchange_rate": _number(totals.get("exchange_rate")),
                "discount_percent": _number(totals.get("discount_percent")),
                "contract_total": _number(totals.get("contract_total")),
                "contract_total_auto": _number(totals.get("contract_total_auto")),
                "rub_total_service_text": _text(totals.get("rub_total_service_text"), 500),
                "discounted_total_service_text": _text(totals.get("discounted_total_service_text"), 500),
                "travel_expenses_mode": _text(totals.get("travel_expenses_mode"), 20) or "actual",
            },
        }

    return {
        "proposal_label": _text(payload.get("proposal_label"), 255),
        "assets": assets,
        "stages": stages,
        "summary": summary,
    }


class CommercialWorkbookBuilder:
    def __init__(self, snapshot):
        self.snapshot = normalize_snapshot(snapshot)
        self.workbook = Workbook()
        self.sheet = self.workbook.active
        self.sheet.title = "Коммерческое предложение"
        self.row = 1
        self.stage_layouts = []
        self.thin_border = Border(
            left=Side(style="thin", color="DEE2E6"),
            right=Side(style="thin", color="DEE2E6"),
            top=Side(style="thin", color="DEE2E6"),
            bottom=Side(style="thin", color="DEE2E6"),
        )
        self.title_fill = PatternFill("solid", fgColor="D9E7F5")
        self.header_fill = PatternFill("solid", fgColor="E8EEF6")
        self.total_fill = PatternFill("solid", fgColor="F3F6FA")
        self.wrap = Alignment(wrap_text=True, vertical="top")

    def _style_row(self, row, first_col, last_col, *, fill=None, bold=False):
        for column in range(first_col, last_col + 1):
            cell = self.sheet.cell(row=row, column=column)
            cell.border = self.thin_border
            cell.alignment = self.wrap
            if fill:
                cell.fill = fill
            if bold:
                cell.font = Font(bold=True)

    def _write_title(self, title, last_col):
        self.sheet.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=last_col)
        cell = self.sheet.cell(self.row, 1, _excel_text(title))
        cell.font = Font(bold=True, size=11)
        cell.fill = self.title_fill
        cell.alignment = self.wrap
        self._style_row(self.row, 1, last_col, fill=self.title_fill, bold=True)
        self.row += 1

    def _write_headers(self, day_headers):
        headers = [
            "Специалист",
            "Специальность",
            "Профессиональный статус",
            "Код",
            "Услуги",
            "Ставка, евро / день",
            *day_headers,
            "Всего",
            "Итого, евро без НДС",
        ]
        for column, value in enumerate(headers, 1):
            self.sheet.cell(self.row, column, value)
        self._style_row(self.row, 1, len(headers), fill=self.header_fill, bold=True)
        self.row += 1
        return {
            "rate_col": 6,
            "day_cols": list(range(7, 7 + len(day_headers))),
            "total_days_col": 7 + len(day_headers),
            "total_col": 8 + len(day_headers),
            "last_col": len(headers),
        }

    def _write_common_values(self, excel_row, row_data):
        values = [
            row_data["specialist"],
            row_data["job_title"],
            row_data["professional_status"],
            row_data["code"],
            row_data["service_name"],
            row_data["rate_eur_per_day"],
        ]
        for column, value in enumerate(values, 1):
            self.sheet.cell(excel_row, column, _excel_text(value) if column < 6 else value)

    def _write_financial_rows(self, columns, totals, data_rows, travel_layout):
        total_col_letter = get_column_letter(columns["total_col"])
        day_letters = [get_column_letter(column) for column in columns["day_cols"]]
        data_excel_rows = [item["excel_row"] for item in data_rows]

        summary_row = self.row
        self.sheet.cell(summary_row, 5, "ИТОГО, по расчёту")
        for column, letter in zip(columns["day_cols"], day_letters):
            self.sheet.cell(summary_row, column, _formula_nonzero_sum(f"{letter}{r}" for r in data_excel_rows))
        day_range = f"{day_letters[0]}{summary_row}:{day_letters[-1]}{summary_row}"
        self.sheet.cell(summary_row, columns["total_days_col"], f'=IF(SUM({day_range})=0,"",SUM({day_range}))')
        total_refs = [f"{total_col_letter}{r}" for r in data_excel_rows]
        self.sheet.cell(summary_row, columns["total_col"], _formula_nonzero_sum(total_refs))
        self._style_row(summary_row, 1, columns["last_col"], fill=self.total_fill, bold=True)
        self.row += 1

        with_travel_row = self.row
        self.sheet.cell(with_travel_row, 5, "ИТОГО, евро с командировочными по расчёту")
        self.sheet.cell(
            with_travel_row,
            columns["total_col"],
            f'=IF(SUM({total_col_letter}{summary_row},{total_col_letter}{travel_layout["excel_row"]})=0,"",'
            f'SUM({total_col_letter}{summary_row},{total_col_letter}{travel_layout["excel_row"]}))',
        )
        self._style_row(with_travel_row, 1, columns["last_col"], fill=self.total_fill, bold=True)
        self.row += 1

        rub_row = self.row
        self.sheet.cell(rub_row, 5, _excel_text(totals["rub_total_service_text"] or "Курс евро Банка России:"))
        self.sheet.cell(rub_row, columns["rate_col"], totals["exchange_rate"])
        self.sheet.cell(
            rub_row,
            columns["total_col"],
            f'=IF(OR({total_col_letter}{with_travel_row}="",{get_column_letter(columns["rate_col"])}{rub_row}=""),"",'
            f'{total_col_letter}{with_travel_row}*{get_column_letter(columns["rate_col"])}{rub_row})',
        )
        self._style_row(rub_row, 1, columns["last_col"], fill=self.total_fill, bold=True)
        self.row += 1

        discounted_row = self.row
        self.sheet.cell(discounted_row, 5, _excel_text(totals["discounted_total_service_text"] or "Размер скидки:"))
        self.sheet.cell(discounted_row, columns["rate_col"], totals["discount_percent"] if totals["discount_percent"] is not None else 5)
        self.sheet.cell(
            discounted_row,
            columns["total_col"],
            f'=IF({total_col_letter}{rub_row}="","",{total_col_letter}{rub_row}*'
            f'(1-{get_column_letter(columns["rate_col"])}{discounted_row}/100))',
        )
        self._style_row(discounted_row, 1, columns["last_col"], fill=self.total_fill, bold=True)
        self.row += 1

        contract_row = self.row
        self.sheet.cell(contract_row, 5, "ИТОГО в договор, рубли без НДС с учётом дополнительной скидки")
        contract_total = totals["contract_total"]
        contract_auto = totals["contract_total_auto"]
        is_manual = contract_total is not None and (
            contract_auto is None or Decimal(str(contract_total)) != Decimal(str(contract_auto))
        )
        if is_manual:
            contract_value = contract_total
        else:
            contract_value = (
                f'=IF({total_col_letter}{discounted_row}="","",'
                f'ROUNDDOWN({total_col_letter}{discounted_row}/100000,0)*100000)'
            )
        self.sheet.cell(contract_row, columns["total_col"], contract_value)
        self._style_row(contract_row, 1, columns["last_col"], fill=self.total_fill, bold=True)
        self.row += 1
        return {
            "summary_row": summary_row,
            "with_travel_row": with_travel_row,
            "rub_row": rub_row,
            "discounted_row": discounted_row,
            "contract_row": contract_row,
        }

    def _write_stage(self, stage, stage_index):
        assets = self.snapshot["assets"]
        columns_count = 8 + len(assets)
        self._write_title(stage["label"], columns_count)
        columns = self._write_headers(assets)
        data_rows = []
        travel_data = next((item for item in stage["rows"] if _is_travel(item)), None)

        for row_data in (item for item in stage["rows"] if not _is_travel(item)):
            excel_row = self.row
            self._write_common_values(excel_row, row_data)
            for column, value in zip(columns["day_cols"], row_data["asset_day_counts"]):
                self.sheet.cell(excel_row, column, value)
            day_start = get_column_letter(columns["day_cols"][0])
            day_end = get_column_letter(columns["day_cols"][-1])
            rate = get_column_letter(columns["rate_col"])
            day_total = get_column_letter(columns["total_days_col"])
            self.sheet.cell(excel_row, columns["total_days_col"], f'=IF(SUM({day_start}{excel_row}:{day_end}{excel_row})=0,"",SUM({day_start}{excel_row}:{day_end}{excel_row}))')
            self.sheet.cell(excel_row, columns["total_col"], f'=IF(OR({rate}{excel_row}="",{day_total}{excel_row}=""),"",{rate}{excel_row}*{day_total}{excel_row})')
            self._style_row(excel_row, 1, columns["last_col"])
            data_rows.append({"data": row_data, "excel_row": excel_row})
            self.row += 1

        travel_data = travel_data or {
            "specialist": "",
            "job_title": "",
            "professional_status": "",
            "code": "",
            "service_name": "Командировочные расходы, евро",
            "rate_eur_per_day": None,
            "asset_day_counts": [None] * len(assets),
            "total_eur_without_vat": None,
        }
        travel_row = self.row
        self._write_common_values(travel_row, travel_data)
        mode = stage["totals"]["travel_expenses_mode"]
        if mode == "calculation":
            for column, value in zip(columns["day_cols"], travel_data["asset_day_counts"]):
                self.sheet.cell(travel_row, column, value)
            day_start = get_column_letter(columns["day_cols"][0])
            day_end = get_column_letter(columns["day_cols"][-1])
            sum_formula = f"SUM({day_start}{travel_row}:{day_end}{travel_row})"
            self.sheet.cell(travel_row, columns["total_days_col"], f'=IF({sum_formula}=0,"",{sum_formula})')
            self.sheet.cell(travel_row, columns["total_col"], f'=IF({sum_formula}=0,"",{sum_formula})')
        else:
            self.sheet.cell(travel_row, columns["rate_col"], "по факту")
            self.sheet.cell(travel_row, columns["total_col"], travel_data["total_eur_without_vat"])
        self._style_row(travel_row, 1, columns["last_col"])
        self.row += 1

        fixed_rows = self._write_financial_rows(
            columns,
            stage["totals"],
            data_rows,
            {"data": travel_data, "excel_row": travel_row},
        )
        layout = {
            "stage_index": stage_index,
            "columns": columns,
            "data_rows": data_rows,
            "travel": {"data": travel_data, "excel_row": travel_row},
            **fixed_rows,
        }
        self.stage_layouts.append(layout)
        self.row += 2

    def _summary_matches(self, summary_row):
        key = _group_key(summary_row)
        matches = []
        for layout in self.stage_layouts:
            for item in layout["data_rows"]:
                if _group_key(item["data"]) == key:
                    matches.append((layout, item))
        return matches

    def _write_summary(self, summary):
        assets = self.snapshot["assets"]
        day_specs = [
            (stage_index, asset_index)
            for stage_index in range(len(self.stage_layouts))
            for asset_index in range(len(assets))
        ]
        day_headers = []
        for stage_index, asset_index in day_specs:
            stage_label = self.snapshot["stages"][stage_index]["product_label"] or f"Этап {stage_index + 1}"
            day_headers.append(f"{stage_label}: {assets[asset_index]}")
        columns_count = 8 + len(day_headers)
        self._write_title(summary["label"], columns_count)
        columns = self._write_headers(day_headers)
        data_rows = []
        travel_data = next((item for item in summary["rows"] if _is_travel(item)), None)

        for row_data in (item for item in summary["rows"] if not _is_travel(item)):
            excel_row = self.row
            self._write_common_values(excel_row, row_data)
            matches = self._summary_matches(row_data)
            if matches:
                first_layout, first_item = matches[0]
                first_rate_cell = (
                    f'{get_column_letter(first_layout["columns"]["rate_col"])}{first_item["excel_row"]}'
                )
                self.sheet.cell(excel_row, columns["rate_col"], f"={first_rate_cell}")
            for target_column, (stage_index, asset_index) in zip(columns["day_cols"], day_specs):
                refs = []
                for layout, item in matches:
                    if layout["stage_index"] != stage_index:
                        continue
                    source_column = layout["columns"]["day_cols"][asset_index]
                    refs.append(f"{get_column_letter(source_column)}{item['excel_row']}")
                self.sheet.cell(excel_row, target_column, _formula_nonzero_sum(refs))
            first_day = get_column_letter(columns["day_cols"][0])
            last_day = get_column_letter(columns["day_cols"][-1])
            rate_col = get_column_letter(columns["rate_col"])
            total_days_col = get_column_letter(columns["total_days_col"])
            self.sheet.cell(excel_row, columns["total_days_col"], f'=IF(SUM({first_day}{excel_row}:{last_day}{excel_row})=0,"",SUM({first_day}{excel_row}:{last_day}{excel_row}))')
            self.sheet.cell(excel_row, columns["total_col"], f'=IF(OR({rate_col}{excel_row}="",{total_days_col}{excel_row}=""),"",{rate_col}{excel_row}*{total_days_col}{excel_row})')
            self._style_row(excel_row, 1, columns["last_col"])
            data_rows.append({"data": row_data, "excel_row": excel_row})
            self.row += 1

        travel_data = travel_data or {
            "specialist": "",
            "job_title": "",
            "professional_status": "",
            "code": "",
            "service_name": "Командировочные расходы, евро",
            "rate_eur_per_day": None,
            "asset_day_counts": [None] * len(assets),
            "total_eur_without_vat": None,
        }
        travel_row = self.row
        self._write_common_values(travel_row, travel_data)
        travel_total_refs = []
        for target_column, (stage_index, asset_index) in zip(columns["day_cols"], day_specs):
            source_layout = self.stage_layouts[stage_index]
            source_day_column = source_layout["columns"]["day_cols"][asset_index]
            source_ref = f"{get_column_letter(source_day_column)}{source_layout['travel']['excel_row']}"
            if summary["totals"]["travel_expenses_mode"] == "calculation":
                self.sheet.cell(travel_row, target_column, f"={source_ref}")
            travel_total_refs.append(
                f'{get_column_letter(source_layout["columns"]["total_col"])}{source_layout["travel"]["excel_row"]}'
            )
        first_day = get_column_letter(columns["day_cols"][0])
        last_day = get_column_letter(columns["day_cols"][-1])
        if summary["totals"]["travel_expenses_mode"] == "calculation":
            self.sheet.cell(travel_row, columns["total_days_col"], f'=IF(SUM({first_day}{travel_row}:{last_day}{travel_row})=0,"",SUM({first_day}{travel_row}:{last_day}{travel_row}))')
        else:
            self.sheet.cell(travel_row, columns["rate_col"], "по факту")
        self.sheet.cell(travel_row, columns["total_col"], _formula_nonzero_sum(travel_total_refs))
        self._style_row(travel_row, 1, columns["last_col"])
        self.row += 1

        self._write_financial_rows(
            columns,
            summary["totals"],
            data_rows,
            {"data": travel_data, "excel_row": travel_row},
        )

    def _finish(self):
        self.sheet.freeze_panes = "A3"
        widths = {1: 20, 2: 24, 3: 24, 4: 12, 5: 42, 6: 20}
        for column, width in widths.items():
            self.sheet.column_dimensions[get_column_letter(column)].width = width
        for column in range(7, self.sheet.max_column + 1):
            self.sheet.column_dimensions[get_column_letter(column)].width = 18
        for row in self.sheet.iter_rows():
            for cell in row:
                if cell.column >= 6 and (isinstance(cell.value, (int, float)) or (isinstance(cell.value, str) and cell.value.startswith("="))):
                    cell.number_format = '#,##0.00'
        self.sheet.auto_filter.ref = None
        self.workbook.calculation.calcMode = "auto"
        self.workbook.calculation.fullCalcOnLoad = True
        self.workbook.calculation.forceFullCalc = True

    def build(self):
        for index, stage in enumerate(self.snapshot["stages"]):
            self._write_stage(stage, index)
        if self.snapshot["summary"]:
            self._write_summary(self.snapshot["summary"])
        self._finish()
        return self.workbook


def build_commercial_xlsx(snapshot):
    workbook = CommercialWorkbookBuilder(snapshot).build()
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
